"""Generic Excel reader — auto-detects firma & phone columns across any sheet structure."""
import re
import pandas as pd
from openpyxl import load_workbook
from models import Contact

# Red fill variations (different Excel apps use different alpha/format)
RED_FILLS = {"FFFF0000", "00FF0000", "FF0000", "FFFF6666", "FFFF3333"}

# Column patterns — (pattern, score). Higher score wins on conflict. Word-boundary matched.
COLUMN_PATTERNS = {
    "firma": [
        ("firma ismi", 100), ("firma adi", 100), ("firma adı", 100),
        ("sirket adi", 90), ("isletme adi", 90), ("magaza adi", 90),
        ("firma", 70), ("sirket", 70), ("isletme", 70), ("kurum", 65),
        ("magaza", 65), ("marka", 60), ("company", 80), ("name", 55),
    ],
    "telefon": [
        ("telefon numarasi", 100), ("telefon no", 95), ("cep telefonu", 95),
        ("iletisim", 90), ("telefon", 85), ("phone", 85), ("gsm", 80),
        ("mobil", 70), ("cep", 75), ("numara", 60), ("tel", 50),
    ],
    "ilgili_kisi": [
        ("ilgili kisi", 100), ("yetkili kisi", 100),
        ("ilgili", 80), ("yetkili", 80), ("muhatap", 75), ("contact person", 90),
    ],
    "departman": [
        ("departman", 90), ("department", 90), ("bolum", 75), ("pozisyon", 70), ("title", 60),
    ],
    "lokasyon": [
        ("lokasyon", 90), ("location", 90), ("sehir", 80), ("city", 80),
        ("ilce", 70), ("adres", 75), ("address", 75), ("konum", 75),
    ],
    "email": [
        ("e-mail", 95), ("e-posta", 95), ("eposta", 90), ("email", 90), ("mail", 70),
    ],
}

# Tokens we never want as firma values (header fragments)
HEADER_TOKENS = {"firma ismi", "firma adi", "firma", "iletisim", "telefon", "ilan", "header", ""}


def _normalize(text: str) -> str:
    """Lowercase + strip TR accents + remove combining marks for fuzzy matching."""
    import unicodedata
    if not isinstance(text, str):
        return ""
    # Replace TR-specific chars first (before lowercase to avoid İ→i+combining)
    s = (text
        .replace("İ", "i").replace("Ş", "s").replace("Ğ", "g")
        .replace("Ü", "u").replace("Ö", "o").replace("Ç", "c")
        .replace("ı", "i").replace("ş", "s").replace("ğ", "g")
        .replace("ü", "u").replace("ö", "o").replace("ç", "c")
    ).lower()
    # Strip any remaining combining marks
    s = "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))
    return s.strip()


def _match_score(col_norm: str, patterns: list) -> tuple:
    """Return (best_score, matched_pattern) for a column against pattern list."""
    if not col_norm:
        return (0, None)
    # Word-boundary match: pattern must appear as a whole token or token sequence
    best = (0, None)
    for pat, score in patterns:
        # Use word boundaries; pat may be multi-word
        pat_re = r"(?:^|\W)" + re.escape(pat) + r"(?:$|\W)"
        if re.search(pat_re, col_norm):
            if score > best[0]:
                best = (score, pat)
    return best


def _detect_columns(df: pd.DataFrame) -> dict:
    """Map our standard field names to actual columns using best-match scoring."""
    found = {}
    used_cols = set()
    cols = list(df.columns)
    norm_cols = [_normalize(str(c)) for c in cols]

    # For each field, pick the column with the highest matching score
    # Process in priority order: firma & telefon first (most important), then others
    field_order = ["firma", "telefon", "ilgili_kisi", "email", "lokasyon", "departman"]

    for field in field_order:
        patterns = COLUMN_PATTERNS[field]
        best_score = 0
        best_col = None
        for i, ncol in enumerate(norm_cols):
            if cols[i] in used_cols:
                continue
            score, _ = _match_score(ncol, patterns)
            if score > best_score:
                best_score = score
                best_col = cols[i]
        if best_col and best_score >= 50:
            found[field] = best_col
            used_cols.add(best_col)

    return found


def _detect_header_row(file_path: str, sheet_name: str, max_scan: int = 8) -> int:
    """Find the row that has BOTH firma and telefon header keywords."""
    raw = pd.read_excel(file_path, sheet_name=sheet_name, header=None, nrows=max_scan, dtype=str).fillna("")
    best_row = 0
    best_score = 0
    for idx, row in raw.iterrows():
        cells = [_normalize(str(c)) for c in row]
        firma_score = max((_match_score(c, COLUMN_PATTERNS["firma"])[0] for c in cells), default=0)
        phone_score = max((_match_score(c, COLUMN_PATTERNS["telefon"])[0] for c in cells), default=0)
        if firma_score >= 50 and phone_score >= 50:
            total = firma_score + phone_score
            if total > best_score:
                best_score = total
                best_row = idx
    return best_row


def _get_red_rows(wb, sheet_name: str) -> set:
    """Find rows where any cell has a red fill."""
    if sheet_name not in wb.sheetnames:
        return set()
    ws = wb[sheet_name]
    red_rows = set()
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            try:
                if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb:
                    if str(cell.fill.fgColor.rgb).upper() in RED_FILLS:
                        red_rows.add(cell.row)
                        break
            except (TypeError, AttributeError):
                pass
    return red_rows


def _is_likely_firma(value: str) -> bool:
    """Filter out garbage / header-like values."""
    if not value:
        return False
    norm = _normalize(value)
    if norm in HEADER_TOKENS:
        return False
    # All digits → probably a phone, not a firma
    digits_only = re.sub(r"\D", "", value)
    if digits_only == value.strip() and len(value.strip()) >= 7:
        return False
    return True


def _is_likely_phone(value: str) -> bool:
    """Has enough digits to plausibly be a phone."""
    if not value:
        return False
    digits = re.sub(r"\D", "", value)
    return len(digits) >= 10


def read_excel(file_path: str) -> list:
    """Read any Excel file, auto-detect firma/phone columns, return Contact list."""
    wb = load_workbook(file_path, read_only=False, data_only=True)
    contacts = []

    for sheet_name in wb.sheetnames:
        try:
            header_row = _detect_header_row(file_path, sheet_name)
            df = pd.read_excel(file_path, sheet_name=sheet_name,
                               header=header_row, dtype=str).fillna("")
        except Exception:
            continue

        if df.empty:
            continue

        # Clean column names
        df.columns = [str(c).strip() for c in df.columns]

        # Auto-detect column mapping
        cmap = _detect_columns(df)

        if "firma" not in cmap or "telefon" not in cmap:
            # Fallback: try first 2 columns as (firma, telefon)
            if len(df.columns) >= 2:
                cmap = {"firma": df.columns[0], "telefon": df.columns[1]}
            else:
                continue

        red_rows = _get_red_rows(wb, sheet_name)

        for idx, row in df.iterrows():
            firma = str(row.get(cmap["firma"], "")).strip()
            telefon = str(row.get(cmap["telefon"], "")).strip()

            # Phone is required, firma can be empty (we use "-")
            if not _is_likely_phone(telefon):
                continue
            # But filter out header-row leftovers (when firma is exactly a header token)
            if firma and _normalize(firma) in HEADER_TOKENS:
                continue
            if not firma:
                firma = "-"

            ilgili_kisi = str(row.get(cmap.get("ilgili_kisi", ""), "")).strip() if cmap.get("ilgili_kisi") else ""
            departman = str(row.get(cmap.get("departman", ""), "")).strip() if cmap.get("departman") else ""
            lokasyon = str(row.get(cmap.get("lokasyon", ""), "")).strip() if cmap.get("lokasyon") else ""
            email = str(row.get(cmap.get("email", ""), "")).strip() if cmap.get("email") else ""

            # Excel row = df index + header offset + 1 (header) + 1 (1-based)
            excel_row = idx + header_row + 2
            is_blacklisted = excel_row in red_rows

            contacts.append(Contact(
                firma=firma,
                telefon=telefon,
                ilgili_kisi=ilgili_kisi,
                departman=departman,
                lokasyon=lokasyon,
                email=email,
                sheet=sheet_name,
                is_blacklisted=is_blacklisted,
            ))

    wb.close()
    return contacts
